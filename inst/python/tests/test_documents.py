"""Readers, chunker, flags, and change detection for document sets. No models, no network."""

import json
import os
import tempfile
from pathlib import Path

from dragonfarm import documents as D

MD = """# Returns

Any device can be returned within 30 days of delivery, opened or unopened, for a full refund.
This includes devices that have already been set up. """ + " ".join(["Filler sentence number %d." % i for i in range(30)]) + """

## How to start

To start a return, open Orders in your account, pick the order, and choose Start a return.
A prepaid label is issued to print or show at any drop-off point. """ + " ".join(["More words %d." % i for i in range(25)]) + """

> We will replace the damaged unit at no cost. Please send a photo of the damage and one of the label.
> Usually a business day after the return is scanned in. """ + " ".join(["Quoted %d." % i for i in range(20)]) + """

## Tiny

Just a few words here.

## Example

```r
x <- dragon_train(ds, "model")
print(x)
```
"""

HTML = """<html><head><title>Help</title><style>p{}</style></head><body>
<nav><a href="/">Home</a> Skip me</nav>
<main>
<h1>Wi-Fi drops</h1>
<p>Our devices connect on the 2.4 GHz band. Mesh routers that use band steering can push a device onto the wrong band,
which causes repeated disconnects and failed setup at the Wi-Fi step. """ + " ".join(["Detail %d." % i for i in range(30)]) + """</p>
<h2>Steps</h2>
<ul><li>Create a dedicated 2.4 GHz network.</li><li>Hold reset for 10 seconds until the light blinks amber.</li></ul>
<table><tr><th>Model</th><th>Memory</th></tr><tr><td>SmolLM2-135M</td><td>2 GB</td></tr></table>
<pre>library(dragonfarm)
dragon_check()</pre>
</main><footer>Built with pkgdown</footer></body></html>"""


def test_markdown_chunks_keep_sections_and_mark_quotes_and_code():
    chunks = D.chunk_text(MD, "policy.md", max_words=120, min_words=20)
    sections = [c["section"] for c in chunks]
    assert "Returns" in sections and "How to start" in sections
    assert all(c["source"] == "policy.md" for c in chunks)
    assert all(c["n_words"] <= 180 for c in chunks)
    quoted = [c for c in chunks if c["quoted"]]
    assert quoted and "no cost" in quoted[0]["text"]
    code = [c for c in chunks if c["code"]]
    assert code and "dragon_train" in code[0]["text"] and code[0]["section"] == "Example"
    # the tiny section was folded into its neighbour, not emitted alone or dropped
    assert not any(c["section"] == "Tiny" for c in chunks)
    assert any("Tiny. Just a few words here." in c["text"] for c in chunks)


def test_a_chunk_never_crosses_a_heading():
    chunks = D.chunk_text(MD, "policy.md")
    for c in chunks:
        assert "# " not in c["text"]


def test_html_reader_skips_chrome_and_keeps_structure():
    text = D.html_to_text(HTML)
    assert "Skip me" not in text and "Built with pkgdown" not in text
    assert "# Wi-Fi drops" in text and "## Steps" in text
    assert "SmolLM2-135M | 2 GB" in text
    assert "```\nlibrary(dragonfarm)" in text
    chunks = D.chunk_text(text, "help.html", min_words=10)
    assert any(c["section"] == "Wi-Fi drops" for c in chunks)


def test_csv_rows_become_column_value_lines():
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "faq.csv"
        p.write_text("question,answer\nHow long is the warranty?,Two years from delivery.\n,\nCan I return it?,Within 30 days.\n",
                     encoding="utf-8")
        text = D.read_csv(str(p))
    assert text.startswith("## faq")
    assert "question: How long is the warranty?; answer: Two years from delivery." in text
    assert text.count("\n\n") == 2  # heading + two rows; the empty row is dropped


def test_changelog_files_are_flagged_by_name_or_by_version_headings():
    assert D.changelog_like("NEWS.md", [])
    assert D.changelog_like("CHANGELOG.txt", [])
    assert not D.changelog_like("guide.md", [{"section": "Install"}, {"section": "Use"}])
    versioned = [{"section": s} for s in ["1.2.0", "1.1.0", "Version 1.0.0", "Intro", "0.9.0"]]
    assert D.changelog_like("history-of-releases-but-named-oddly.md", versioned)


def test_build_writes_manifest_and_reuses_unchanged_sources():
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "src"
        src.mkdir()
        (src / "policy.md").write_text(MD, encoding="utf-8")
        (src / "help.html").write_text(HTML, encoding="utf-8")
        (src / "NEWS.md").write_text("# 1.1.0\n\n" + " ".join(["Changed thing %d." % i for i in range(60)]), encoding="utf-8")
        (src / "notes.xyz").write_text("ignored", encoding="utf-8")
        out = Path(tmp) / "set"
        s1 = D.build([str(src), str(src / "missing.md")], out, "test", max_words=120, min_words=20)
        assert s1["n_sources"] == 4 and s1["n_failed"] == 1 and s1["failed"][0]["name"] == "missing.md"
        assert s1["n_chunks"] > 0 and s1["n_reused"] == 0
        manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
        by_name = {s["name"]: s for s in manifest["sources"]}
        assert by_name["NEWS.md"]["changelog_like"] is True and by_name["policy.md"]["changelog_like"] is False
        chunks = D.load_chunks(out)
        assert all(c["id"].startswith(c["source_id"]) for c in chunks)
        news = [c for c in chunks if c["source"] == "NEWS.md"]
        assert news and all(c["weight"] == 0.7 for c in news)
        assert all(c["weight"] == 1.0 for c in chunks if c["source"] == "policy.md")

        # Second build: nothing changed, every readable source is reused and ids are stable.
        s2 = D.build([str(src)], out, "test", max_words=120, min_words=20)
        assert s2["n_reused"] == 3 and s2["n_chunks"] == s1["n_chunks"]
        assert [c["id"] for c in D.load_chunks(out)] == [c["id"] for c in chunks]

        # Change one file: only it is re-read, and its chunk ids change.
        (src / "policy.md").write_text(MD + "\n\n## Extra\n\n" + " ".join(["Added %d." % i for i in range(40)]), encoding="utf-8")
        s3 = D.build([str(src)], out, "test", max_words=120, min_words=20)
        assert s3["n_reused"] == 2
        ids3 = {c["source"]: c["id"] for c in D.load_chunks(out)}
        assert ids3["help.html"] == {c["source"]: c["id"] for c in chunks}["help.html"]
        assert ids3["policy.md"] != {c["source"]: c["id"] for c in chunks}["policy.md"]

        # Different chunking parameters invalidate the reuse.
        s4 = D.build([str(src)], out, "test", max_words=60, min_words=10)
        assert s4["n_reused"] == 0


def test_unsupported_and_missing_files_are_reported_not_raised():
    with tempfile.TemporaryDirectory() as tmp:
        bad = Path(tmp) / "x.bin"
        bad.write_bytes(b"\x00\x01")
        s = D.build([str(bad), str(Path(tmp) / "nope.md")], Path(tmp) / "set", "t")
        assert s["n_failed"] == 2 and s["n_chunks"] == 0
        errors = " ".join(f["error"] for f in s["failed"])
        assert "unsupported file type" in errors and "does not exist" in errors


def test_pdf_and_docx_readers_when_libraries_are_present():
    try:
        import pypdf  # noqa: F401
    except ImportError:
        pypdf = None
    try:
        import docx  # noqa: F401
    except ImportError:
        docx = None
    with tempfile.TemporaryDirectory() as tmp:
        if docx is not None:
            d = docx.Document()
            d.add_heading("Warranty", level=1)
            d.add_paragraph("Every device carries a 2-year limited warranty counted from the delivery date. " * 6)
            d.add_heading("Claims", level=2)
            d.add_paragraph("Ask for the order number or the serial number on the base of the unit. " * 6)
            p = Path(tmp) / "w.docx"
            d.save(str(p))
            text = D.read_docx(str(p))
            assert "# Warranty" in text and "## Claims" in text and "serial number" in text
        if pypdf is not None:
            w = pypdf.PdfWriter()
            w.add_blank_page(width=300, height=300)
            p = Path(tmp) / "blank.pdf"
            with open(p, "wb") as fh:
                w.write(fh)
            try:
                D.read_pdf(str(p))
                raise AssertionError("an image-only PDF should be reported, not read as empty")
            except ValueError as e:
                assert "scanned PDF" in str(e) and "OCR" in str(e)


def test_a_tiny_source_still_yields_one_chunk():
    text = D.rows_to_text("faq", ["question", "answer"], [["How long is the warranty?", "Two years."]])
    chunks = D.chunk_text(text, "faq.csv")
    assert len(chunks) == 1 and "Two years." in chunks[0]["text"] and chunks[0]["section"] == "faq"
    assert D.chunk_text("", "empty.txt") == []


def test_private_use_glyphs_are_stripped():
    chunks = D.chunk_text("# T\n\n Icon " + " ".join(["word"] * 40), "a.md", min_words=10)
    assert "" not in chunks[0]["text"]


def test_key_value_sheets_are_not_read_as_header_plus_rows():
    rows = [("Source", "R datasets package, state.x77 (US Department of Commerce, Statistical Abstract, 1977)"),
            ("Population", "July 1, 1975 estimate, in thousands")]
    assert D.key_value_sheet(rows)
    assert not D.looks_like_header(rows[0], rows[1:])
    text = D.rows_to_text("Notes", None, rows)
    assert "Source: R datasets package" in text and "Population: July 1, 1975" in text


def test_ordinary_tables_keep_their_header():
    rows = [("State", "Region", "Population"), ("Alaska", "West", 365), ("Texas", "South", 12237)]
    assert D.looks_like_header(rows[0], rows[1:]) and not D.key_value_sheet(rows)
    assert not D.looks_like_header((2024, 2023, 2022), rows[1:])     # numbers are data, not names


def test_sheets_never_share_a_chunk():
    big = D.rows_to_text("Data", ["State", "Population"], [("State %d" % i, i * 1000) for i in range(30)])
    notes = D.rows_to_text("Notes", None, [("Source", "Census estimates")])
    text = big + "\n\n" + D.HARD_BREAK + "\n\n" + notes
    chunks = D.chunk_text(text, "book.xlsx", max_words=170, min_words=35)
    last = chunks[-1]
    assert last["section"] == "Notes" and "Source: Census estimates" in last["text"]
    assert all("Census" not in c["text"] for c in chunks[:-1])
    assert all(D.HARD_BREAK not in c["text"] for c in chunks)


def _tiny_pdf(content, resources):
    """A one-page PDF built by hand: content stream and resources as given."""
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Contents 4 0 R /Resources " + resources + b" >>",
            b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceGray "
            b"/BitsPerComponent 8 /Length 1 >>\nstream\n\x80\nendstream"]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)


def test_a_short_pdf_is_read_and_only_a_picture_pdf_is_called_scanned():
    try:
        import pypdf  # noqa: F401
    except ImportError:
        return
    with tempfile.TemporaryDirectory() as tmp:
        short = Path(tmp) / "notice.pdf"
        short.write_bytes(_tiny_pdf(b"BT /F1 12 Tf 20 150 Td (The office is closed on Friday.) Tj ET",
                                    b"<< /Font << /F1 5 0 R >> >>"))
        assert "closed on Friday" in D.read_pdf(str(short))
        scan = Path(tmp) / "scan.pdf"
        scan.write_bytes(_tiny_pdf(b"q 300 0 0 300 0 0 cm /Im1 Do Q BT /F1 8 Tf 5 5 Td (Page 1) Tj ET",
                                   b"<< /Font << /F1 5 0 R >> /XObject << /Im1 6 0 R >> >>"))
        try:
            D.read_pdf(str(scan))
            raise AssertionError("a picture of a page with only a stamp of text should be reported as scanned")
        except ValueError as e:
            assert "scanned PDF" in str(e)
